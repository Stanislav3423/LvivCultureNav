# Генерує звіт ЛР3 (Методології_ЛР3.docx у корені репозиторію).
# Потрібен python-docx:  pip install python-docx  →  python docs/lr3/build_report.py
# Лістинги «до» беруться з тегу v0.2, «після» — з останнього коміту, що змінював код; знімки — з docs/lr3/img.
# Знімки Trello (trello-start.png, trello-done.png) вставляються, якщо вони вже є в docs/lr3/img.
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
IMG = ROOT / 'docs' / 'lr3' / 'img'
OUT = ROOT / 'Методології_ЛР3.docx'
FONT = 'Times New Roman'
MONO = 'Courier New'

doc = Document()
counters = {'table': {}, 'fig': {}, 'listing': {}}
section_no = 0


# ---------- стилі ----------
def setup_styles():
    st = doc.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    pf = st.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    for name, size in (('Heading 1', 14), ('Heading 2', 14)):
        h = doc.styles[name]
        h.font.name = FONT
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.italic = False
        h.font.color.rgb = RGBColor(0, 0, 0)
        rpr = h.element.get_or_add_rPr()
        fonts = rpr.find(qn('w:rFonts'))
        if fonts is None:
            fonts = OxmlElement('w:rFonts')
            rpr.append(fonts)
        for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
            fonts.set(qn(attr), FONT)
        for theme_attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
            if fonts.get(qn(theme_attr)) is not None:
                del fonts.attrib[qn(theme_attr)]
        h.paragraph_format.space_before = Pt(6)
        h.paragraph_format.space_after = Pt(0)
        h.paragraph_format.line_spacing = 1.5
        h.paragraph_format.keep_with_next = True
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.left_margin, sec.right_margin = Cm(2.5), Cm(1.5)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    # Word оновить зміст при відкритті
    settings = doc.settings.element
    upd = OxmlElement('w:updateFields')
    upd.set(qn('w:val'), 'true')
    settings.append(upd)


# ---------- примітиви ----------
def add_runs(par, text, size=None, italic=False):
    """Текст із **жирними** фрагментами і `кодом`."""
    bold = False
    for i, chunk in enumerate(text.split('**')):
        bold = i % 2 == 1
        for j, part in enumerate(chunk.split('`')):
            if not part:
                continue
            run = par.add_run(part)
            run.bold = bold
            run.italic = italic
            if j % 2 == 1:
                run.font.name = MONO
                run.font.size = Pt(12)
            elif size:
                run.font.size = Pt(size)
    return par


def p(text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, size=None, italic=False, space_after=0):
    par = doc.add_paragraph()
    par.alignment = align
    par.paragraph_format.first_line_indent = Cm(1.25) if indent else Cm(0)
    par.paragraph_format.space_after = Pt(space_after)
    return add_runs(par, text, size, italic)


def h1(text):
    par = doc.add_heading(text, level=1)
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return par


def h2(text):
    global section_no
    section_no += 1
    par = doc.add_heading(f'{section_no}. {text}', level=2)
    par.paragraph_format.first_line_indent = Cm(1.25)
    return par


def h3(text):
    par = p(f'**{text}**')
    par.paragraph_format.keep_with_next = True
    return par


def bullets(items):
    for item in items:
        par = doc.add_paragraph(style='List Bullet')
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        par.paragraph_format.line_spacing = 1.5
        add_runs(par, item)


def next_no(kind):
    n = counters[kind].get(section_no, 0) + 1
    counters[kind][section_no] = n
    return f'{section_no}.{n}'


def caption(text, align=WD_ALIGN_PARAGRAPH.LEFT, keep=True):
    par = p(text, align=align, indent=False)
    par.paragraph_format.keep_with_next = keep
    return par


def set_cell_text(cell, text, bold=False, size=11):
    cell.text = ''
    par = cell.paragraphs[0]
    par.paragraph_format.line_spacing = 1.0
    par.paragraph_format.first_line_indent = Cm(0)
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in add_runs(par, f'**{text}**' if bold else text, size=size).runs:
        run.font.size = Pt(size)


def table(title, headers, rows, widths):
    no = next_no('table')
    caption(f'Таблиця {no} - {title}')
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.line_spacing = 1.0
    return no


def figure(file, title, width=16.5):
    no = next_no('fig')
    path = IMG / file
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.first_line_indent = Cm(0)
    par.paragraph_format.keep_with_next = True
    par.paragraph_format.line_spacing = 1.0
    if path.exists():
        par.add_run().add_picture(str(path), width=Cm(width))
    else:
        run = par.add_run(f'[Тут має бути знімок: docs/lr3/img/{file}]')
        run.bold = True
        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    caption(f'Рис. {no} - {title}', align=WD_ALIGN_PARAGRAPH.CENTER, keep=False)
    return no


def git_show(ref, path):
    return subprocess.run(
        ['git', 'show', f'{ref}:{path}'], cwd=ROOT, capture_output=True, check=True, encoding='utf-8'
    ).stdout.splitlines()


def last_code_commit():
    """Код після рефакторингу: останній коміт, що змінював вихідний код (тег v0.3 ставиться пізніше)"""
    return subprocess.run(
        ['git', 'log', '-1', '--format=%H', '--', 'server/src', 'client/src', 'shared/src'],
        cwd=ROOT, capture_output=True, check=True, encoding='utf-8',
    ).stdout.strip()


def fragment(lines, start, end):
    """Рядки від першого, що містить start: до рядка з end (рядок) / end рядків (число) /
    кінця файлу ('EOF') / закриваючої дужки на тому ж відступі (None)"""
    i = next(k for k, line in enumerate(lines) if start in line)
    if end == 'EOF':
        return lines[i:]
    if isinstance(end, int):
        return lines[i : i + end]
    if isinstance(end, str):
        j = next(k for k in range(i, len(lines)) if end in lines[k])
        return lines[i : j + 1]
    indent = len(lines[i]) - len(lines[i].lstrip())
    for j in range(i + 1, len(lines)):
        line = lines[j]
        if line.strip() and len(line) - len(line.lstrip()) == indent and line.strip()[0] in '})]':
            return lines[i : j + 1]
    return lines[i:]


def code(ref, parts):
    """parts: [(шлях, початок, кінець)]; кожна частина підписана шляхом до файлу"""
    out = []
    for path, start, end in parts:
        if out:
            out.append('')
        out.append(f'// {path}')
        out.extend(fragment(git_show(ref, path), start, end))
    return '\n'.join(out)


def listing(title, text):
    no = next_no('listing')
    caption(f'Лістинг {no} - {title}')
    lines = text.split('\n')
    for i, line in enumerate(lines):
        par = doc.add_paragraph()
        par.paragraph_format.line_spacing = 1.0
        par.paragraph_format.first_line_indent = Cm(0)
        par.paragraph_format.keep_with_next = i < len(lines) - 1 and len(lines) < 40
        run = par.add_run(line if line else ' ')
        run.font.name = MONO
        run.font.size = Pt(8)
        run._element.rPr.rFonts.set(qn('w:eastAsia'), MONO)
    doc.add_paragraph().paragraph_format.line_spacing = 1.0
    return no


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_field(par, instr):
    run = par.add_run()
    for tag, text in (('begin', None), (None, instr), ('separate', None), (None, ''), ('end', None)):
        if tag:
            el = OxmlElement('w:fldChar')
            el.set(qn('w:fldCharType'), tag)
        else:
            el = OxmlElement('w:instrText' if text else 'w:t')
            el.set(qn('xml:space'), 'preserve')
            el.text = text or ''
        run._r.append(el)


# ---------- титульна сторінка ----------
def title_page():
    def line(text, align, bold=False, italic=False, space=0):
        par = doc.add_paragraph()
        par.alignment = align
        par.paragraph_format.space_after = Pt(space)
        par.paragraph_format.line_spacing = 1.15
        add_runs(par, f'**{text}**' if bold else text, italic=italic)
        return par

    def blank(n=1):
        for _ in range(n):
            doc.add_paragraph().paragraph_format.line_spacing = 1.0

    C, R = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT
    line('МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ', C, bold=True, space=6)
    line('НАЦІОНАЛЬНИЙ УНІВЕРСИТЕТ "ЛЬВІВСЬКА ПОЛІТЕХНІКА"', C, bold=True)
    blank(2)
    line('Інститут **КНІТ**', R)
    line('Кафедра **ПЗ**', R)
    blank(4)
    line('ЗВІТ', C, bold=True, space=6)
    line('До лабораторної роботи № 3', C)
    par = line('**З дисципліни:** ', C)
    add_runs(par, '“Методології розробки програмного забезпечення”', italic=True)
    par = line('**На тему:** ', C)
    add_runs(par, '“Рефакторинг розробленого коду”', italic=True)
    blank()
    for text, bold in (
        ('Лектор:', True), ('доц. каф. ПЗ', False), ('Сенів М. М.', False), ('', False),
        ('Виконали:', True), ('ст. гр. ПЗІП-12', False), ('Баран В. Б.', False), ('Володькін С. В.', False),
        ('Гавриш Д. В.', False), ('Яскульська Х. В.', False), ('', False),
        ('Прийняв:', True), ('доц. каф. ПЗ', False), ('Федорчук Є.Н.', False),
    ):
        line(text, R, bold=bold)
    blank(2)
    line('« ____ » ________  2026 р.', R)
    line('∑= ____', R)
    blank(2)
    line('Львів – 2026', C)


def footer_page_numbers():
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    # титульна сторінка не нумерується: «Зміст» — сторінка 1, як у звітах ЛР1–ЛР2
    pg = OxmlElement('w:pgNumType')
    pg.set(qn('w:start'), '0')
    sec._sectPr.find(qn('w:cols')).addprevious(pg)
    par = sec.footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_field(par, 'PAGE')




# ---------- зміст звіту ----------
BEFORE = 'v0.2'
AFTER = last_code_commit()

FINDINGS = [
    ('R-01', 'Функція валідації події перевіряє всі поля одним ланцюжком умов; адміністративна форма містить '
             'копії англомовних повідомлень сервера як ключі перекладу',
     'Довга функція (Long Function), «стрілянина дробом» (Shotgun Surgery)',
     'Декомпозиція умови на окремі перевірки полів; винесення текстів помилок у спільний модуль '
     '(Decompose Conditional, Replace Magic Literal)'),
    ('R-02', 'Формування відповіді з помилкою і розбір ідентифікатора повторюються в усіх обробниках; '
             'обробник імпорту партнерських подій містить бізнес-логіку',
     'Дубльований код (Duplicated Code), розбіжні зміни (Divergent Change)',
     'Виділення спільних функцій HTTP-шару; розділення імпорту на етапи перевірки і запису та перенесення '
     'в сервісний шар (Extract Function, Split Phase, Move Function)'),
    ('R-03', 'Сервіс маршрутів одночасно кешує результати, обмежує частоту запитів, звертається до OSRM '
             'і перетворює відповідь; формула округлення часу дубльована',
     'Розбіжні зміни, дубльований код',
     'Виділення кешу та обмежувача частоти в окремі модулі; відокремлення розбору відповіді '
     '(Extract Class, Split Phase)'),
    ('R-04', 'Функція створення репозиторію займає 98 рядків; списки колонок повторюються; заповнення бази '
             'демонстраційними даними продубльоване у двох модулях; рядок «admin» використовується як '
             'магічне значення',
     'Довга функція, дубльований код',
     'Розбиття на короткі іменовані функції; заміна алгоритму групування; іменована константа джерела '
     '(Extract Function, Substitute Algorithm, Replace Magic Literal)'),
    ('R-05', 'Стан форми адміністратора кодується одним значенням типу number | null | undefined з трьома '
             'неявними змістами; обробку помилки авторизації повторено тричі',
     'Одержимість примітивами (Primitive Obsession), дубльований код',
     'Явний тип стану форми; винесення сесії адміністратора в окремий хук '
     '(Replace Primitive with Object, Extract Function)'),
    ('R-06', 'Компонент форми події займає 217 рядків і містить шість однотипних блоків полів, перетворення '
             'даних, мапу і переклад помилок',
     'Великий клас (Large Class), дубльований код',
     'Винесення перетворень даних у чистий модуль; виділення компонентів поля, вкладок мов, мапи і списку '
     'помилок (Extract Function, Extract Class)'),
    ('R-07', 'Параметри запиту маршруту кодуються рядком «термінал:ідентифікатор» і розбираються назад; '
             'асинхронне завантаження даних реалізовано в компоненті сторінки',
     'Одержимість примітивами, довга функція',
     'Заміна рядка об’єктом запиту; винесення завантаження подій і маршруту в окремі хуки '
     '(Replace Primitive with Object, Extract Function)'),
]


def build():
    setup_styles()
    doc.core_properties.author = 'Яскульська Х. В.'
    doc.core_properties.title = 'Звіт до лабораторної роботи № 3'
    title_page()
    footer_page_numbers()
    page_break()

    p('**Зміст**', indent=False)
    toc = doc.add_paragraph()
    add_field(toc, 'TOC \\o "1-2" \\h \\z \\u')
    page_break()

    p('**ТЕМА:** Рефакторинг розробленого коду.', space_after=6)
    p('**МЕТА:** Набути практичних навичок аналізу якості програмного коду, виявлення його недоліків '
      '(«запахів коду») та їх усунення методами рефакторингу без зміни зовнішньої поведінки програми; '
      'навчитися обліковувати недоліки коду як окремі задачі в системі керування проєктами.')

    h1('ЗАВДАННЯ')
    bullets([
        'Ознайомитись з інструкцією та завданням до виконання роботи.',
        'Проаналізувати розроблений в лабораторній роботі № 2 код.',
        'Здійснити рефакторинг розробленого коду.',
        'Завести виявлені недоліки коду як окремі задачі у Trello або Atlassian Jira і відстежувати їх усунення.',
        'У звіті відобразити протокол роботи програми, початковий вихідний код, вихідний код після '
        'рефакторингу, знімки дошок Trello або Jira та зазначити недоліки початкового коду і способи їх усунення.',
    ])

    h1('ХІД ВИКОНАННЯ')

    # 1
    h2('Аналіз програмного коду')
    p('Об’єктом рефакторингу є інформаційно-навігаційна система **LvivCultureNav**, розроблена в лабораторній '
      'роботі № 2. Система побудована як монорепозиторій мовою TypeScript і складається з чотирьох частин: '
      'спільного модуля типів і констант (shared), серверної частини на Express із базою даних SQLite (server), '
      'клієнтського застосунку на React із мапою Leaflet (client) та наскрізних тестів (e2e). Вихідний код '
      'налічує 31 файл; поведінку програми фіксують 38 автоматизованих тестів трьох рівнів.')
    p('Відповідно до визначення М. Фаулера, рефакторинг — це зміна внутрішньої структури програмного '
      'забезпечення, яка полегшує його розуміння і модифікацію, не змінюючи спостережуваної поведінки. '
      'Тому перед початком роботи було перевірено, що всі 38 тестів виконуються успішно: у подальшому вони '
      'слугували критерієм незмінності поведінки після кожного перетворення.')
    p('Аналіз коду проводився методом рецензування в парах. Кожна пара переглядала окремий шар системи — '
      'серверний, клієнтський або спільний — і зіставляла код із каталогом «запахів коду» М. Фаулера: '
      'дубльований код, довга функція, великий клас, розбіжні зміни, «стрілянина дробом», одержимість '
      'примітивами. Для кожного виявленого фрагмента фіксувалися запах коду, наслідки для супроводу програми '
      'та прийом рефакторингу, яким його доцільно усунути. Близькі за змістом зауваження в межах одного '
      'модуля об’єднувалися, щоб кожна задача відповідала цілісній зміні структури коду.')
    p('У результаті аналізу виявлено сім недоліків, які зведено в табл. 1.1. Найпоширенішими виявилися довгі '
      'функції й компоненти з кількома відповідальностями та смислове дублювання; у клієнтській частині, '
      'крім того, трапляється одержимість примітивами — кодування складного стану простими значеннями.')
    table('Недоліки початкового коду і способи їх усунення', ['ID', 'Недолік', 'Запах коду', 'Спосіб усунення'],
          FINDINGS, [1.4, 6.2, 3.6, 5.8])

    # 2
    h2('Облік задач у Trello')
    p('Для обліку робіт на дошці LvivCultureNav створено ітерацію I3 «Рефакторинг» із колонками Backlog, '
      'To Do, In Progress, Testing і Done, як і в попередніх лабораторних роботах. Кожен недолік із табл. 1.1 '
      'оформлено окремою карткою R-01…R-07, що містить опис запаху коду, його розташування, запланований '
      'прийом рефакторингу і критерій готовності. Додаткові картки описують аналіз коду, оцінювання й '
      'пріоритезацію задач (planning game), розроблення тестів для виділених модулів і регресійне '
      'тестування.')
    p('Як і в лабораторній роботі № 2, рефакторинг виконували всі учасники команди відповідно до практики '
      'колективного володіння кодом. Роботу вели в парах згідно з розподілом карток: кожну картку '
      'призначено двом учасникам, один із яких вносив зміни, а другий одразу рецензував код. Картка '
      'переходила в колонку Testing після завершення перетворень і в Done — після успішного проходження '
      'модульних, інтеграційних та E2E-тестів (рис. 2.1, 2.2).')
    figure('trello-start.png', 'Дошка Trello на початку ітерації I3', width=15)
    figure('trello-done.png', 'Дошка Trello після завершення рефакторингу', width=10)

    # 3
    h2('Рефакторинг програмного коду')
    p('Рефакторинг виконувався послідовністю малих перетворень з каталогу М. Фаулера. Після кожного '
      'перетворення запускалися статичний аналізатор ESLint, перевірка типів TypeScript, модульні та '
      'інтеграційні тести, а після змін у клієнтській частині — також E2E-тести. Тести лабораторної роботи № 2 '
      'під час рефакторингу не змінювалися, оскільки саме вони підтверджують незмінність поведінки програми. '
      'Для модулів, виділених у процесі рефакторингу, розроблено 24 нові модульні тести.')

    h3('3.1. Валідація події (R-01)')
    p('Функція `validateEventInput` перевіряла дату, координати, посилання на медіафайл і переклади одним '
      'ланцюжком умов, а належність координат до меж Львова — довгим логічним виразом. Сервер повертав '
      'помилки англомовними рядками, а адміністративна форма зберігала копії цих рядків як ключі для '
      'перекладу українською (лістинг 3.1). Така залежність є прикладом «стрілянини дробом»: зміна тексту '
      'помилки на сервері непомітно порушувала переклад у клієнті.')
    p('Умову декомпоновано на іменовані перевірки, по одній для кожного поля, які об’єднано в масив '
      '`CHECKS`. Перевірку меж Львова перенесено до константи `LVIV_BOUNDS` у спільний модуль, а тексти '
      'помилок оголошено один раз у константі `VALIDATION_ERRORS`, яку використовують і сервер, і клієнт '
      '(лістинг 3.2). Тексти відповідей API при цьому не змінилися.')
    listing('Переклад помилок валідації в адміністративній формі до рефакторингу',
            code(BEFORE, [('client/src/admin/EventForm.tsx', '// Повідомлення сервера', '};')]))
    listing('Спільні тексти помилок і перевірки полів після рефакторингу',
            code(AFTER, [('shared/src/index.ts', 'export const VALIDATION_ERRORS', None),
                         ('server/src/domain/validateEvent.ts', 'const CHECKS', None),
                         ('client/src/admin/formParts.tsx', 'const MESSAGES', 3)]))

    h3('3.2. HTTP-шар та імпорт подій партнерів (R-02)')
    p('Кожен обробник REST API самостійно формував відповідь з помилкою, а ідентифікатор події розбирався '
      'в різних місцях по-різному. Обробник `POST /api/partners/import` поєднував розбір HTTP-запиту з '
      'бізнес-правилами: перевіряв кожну подію, шукав її за зовнішнім ідентифікатором, створював або '
      'оновлював запис і підраховував результат (лістинг 3.3).')
    p('Повторювані фрагменти виділено у функції `sendError`, `notFound` і `parseId`. Імпорт розділено на два '
      'етапи — перевірку події без звернень до бази даних і запис — та перенесено в сервісний модуль '
      '`partnerImport.ts`; обробник лише перевіряє форму тіла запиту (лістинг 3.4). Завдяки цьому логіку '
      'імпорту можна перевіряти модульними тестами без запуску HTTP-сервера.')
    listing('Обробник імпорту подій до рефакторингу (server/src/api/partners.ts)',
            code(BEFORE, [('server/src/api/partners.ts', "router.post('/import'", None)]))
    listing('Обробник і сервіс імпорту після рефакторингу',
            code(AFTER, [('server/src/api/partners.ts', "router.post('/import'", None),
                         ('server/src/services/partnerImport.ts', 'export function importPartnerEvents', None)]))

    h3('3.3. Сервіс маршрутів (R-03)')
    p('Функція `createOsrmRouting` мала чотири незалежні причини для змін: кешування маршрутів, обмеження '
      'частоти запитів відповідно до політики сервісу OSRM, виконання HTTP-запиту і перетворення відповіді. '
      'Кеш і обмежувач частоти виділено в окремі модулі `RouteCache` і `createThrottle`, перетворення '
      'відповіді — у функцію `toRouteResult`, а формулу округлення часу — у функцію `minutesFromSeconds` '
      'доменного модуля. Поведінка сервісу збережена: кеш, як і раніше, витісняє найстаріший запис і не '
      'зберігає запасний маршрут по прямій.')

    h3('3.4. Репозиторій подій і заповнення бази даних (R-04)')
    p('Функцію `createEventsRepo` обсягом 98 рядків розбито на короткі функції підготовки запитів, '
      'перетворення рядків бази даних і очищення перекладів. Групування перекладів, яке на кожному кроці '
      'копіювало масив, замінено одним проходом зі словником. Заповнення бази демонстраційними подіями '
      'зосереджено у функціях `seedIfEmpty` і `reseed`, а рядок «admin» замінено константою `ADMIN_SOURCE`.')

    h3('3.5. Стан сторінки адміністратора (R-05)')
    p('Стан бічної форми кодувався одним значенням `editingId` з трьома неявними змістами: undefined — '
      'форма закрита, null — створення нової події, число — редагування наявної; зрозуміти це можна було '
      'лише з коментаря (лістинг 3.5). Значення замінено явним типом `FormState` (лістинг 3.6), а вхід, '
      'вихід, перезавантаження списку та обробку помилки авторизації зосереджено в хуку `useAdminSession`.')
    listing('Стан форми до рефакторингу (client/src/admin/AdminPage.tsx)',
            code(BEFORE, [('client/src/admin/AdminPage.tsx', '// undefined — форма закрита', 2),
                          ('client/src/admin/AdminPage.tsx', 'const editing = typeof editingId', 1),
                          ('client/src/admin/AdminPage.tsx', '{(editingId === null || editing) && (', 1)]))
    listing('Стан форми після рефакторингу (client/src/admin/AdminPage.tsx)',
            code(AFTER, [('client/src/admin/AdminPage.tsx', 'type FormState', 1),
                         ('client/src/admin/AdminPage.tsx', 'const editingId = form.mode', 2),
                         ('client/src/admin/AdminPage.tsx', "{(form.mode === 'new' || editing) && (", 1)]))

    h3('3.6. Форма події (R-06)')
    p('Компонент `EventForm` обсягом 217 рядків містив шість однотипних блоків «підпис + поле введення», '
      'перетворення даних форми в тіло запиту, мапу вибору точки та переклад помилок сервера. Перетворення '
      'даних винесено в чистий модуль `eventDraft.ts`, однотипну розмітку — у компоненти `TextField`, '
      '`LangTabs` і `ErrorList`, мапу — у компонент `PointPicker`. Адресу тайлів OpenStreetMap і координати '
      'центру мапи, що повторювалися в двох компонентах, винесено в модуль `map.ts`. Обсяг компонента форми '
      'зменшився до 112 рядків.')

    h3('3.7. Маршрут на терміналі (R-07)')
    p('Компонент `TerminalPage` поєднував ідентифікатор терміналу і події в рядок вигляду `station:5`, а '
      'потім розбирав його назад за допомогою `split` і приведення типу (лістинг 3.7). Рядок замінено '
      'об’єктом `{ terminal, eventId }`: хук `useRoute` зберігає відповідь сервера разом із параметрами '
      'запиту і визначає застарілу відповідь порівнянням полів (лістинг 3.8). Завантаження подій винесено '
      'в хук `useEvents`, завдяки чому обсяг компонента зменшився зі 112 до 63 рядків.')
    listing('Запит маршруту до рефакторингу (client/src/terminal/TerminalPage.tsx)',
            code(BEFORE, [('client/src/terminal/TerminalPage.tsx', '// Маршрут перебудовується', 20)]))
    listing('Хук useRoute після рефакторингу',
            code(AFTER, [('client/src/terminal/hooks.ts', 'export function useRoute', None),
                         ('client/src/terminal/TerminalPage.tsx', 'const route = useRoute', 1)]))

    # 4
    h2('Протокол роботи програми')
    p('Після рефакторингу програму було зібрано й запущено з порожньою базою даних, яку сервер заповнив '
      'демонстраційними подіями. Зовнішня поведінка програми не змінилася. Термінал на вокзалі відображає '
      'події на мапі та будує пішохідний маршрут вулицями (рис. 4.1). За недоступності сервісу маршрутів для '
      'польської мови подію без польського перекладу показано англійською з відповідною позначкою, а маршрут — '
      'пунктирною прямою з орієнтовною відстанню (рис. 4.2).')
    figure('terminal-station.png', 'Термінал на вокзалі: маршрут до Музею Пінзеля', width=15)
    figure('terminal-center-pl.png', 'Термінал у центрі міста: польська мова, запасний переклад', width=11)
    p('Сторінка адміністратора відображає таблицю подій зі станом перекладів і форму редагування (рис. 4.3); '
      'помилки валідації, тексти яких тепер задано спільними константами, перекладаються українською '
      '(рис. 4.4).')
    figure('admin.png', 'Сторінка адміністратора: таблиця подій і редагування перекладу', width=13)
    figure('admin-errors.png', 'Повідомлення про помилки валідації', width=8)
    p('Протокол звернень до REST API (рис. 4.5) підтверджує, що коди й тексти відповідей сервера не '
      'змінилися: повернення запасної мови, помилки 404 і 400 для некоректних параметрів, пакетний імпорт '
      'із відхиленням некоректної події, перелік помилок валідації та відмова без ключа адміністратора.')
    figure('log-api.png', 'Протокол роботи сервера та виклики REST API')
    p('Регресійне тестування показало, що всі 62 тести виконуються успішно, зокрема 38 тестів лабораторної '
      'роботи № 2, які під час рефакторингу не змінювалися (рис. 4.6).')
    figure('log-integration.png', 'Результат інтеграційних тестів REST API')

    h1('ВИСНОВКИ')
    p('У ході лабораторної роботи проаналізовано програмний код системи LvivCultureNav, розробленої в '
      'лабораторній роботі № 2, і виявлено сім недоліків. Найпоширенішими з них є довгі функції та '
      'компоненти з кількома відповідальностями, смислове дублювання коду та одержимість примітивами. '
      'Кожен недолік зареєстровано як окрему задачу на дошці Trello і відстежено до його усунення.')
    p('Недоліки усунуто прийомами рефакторингу з каталогу М. Фаулера: виділенням функцій, класів і '
      'компонентів, декомпозицією умов, розділенням обробки на етапи, заміною магічних значень іменованими '
      'константами та примітивів — явними типами. У результаті кожен модуль має одну причину для змін, '
      'тексти помилок валідації визначено в одному місці, а стан сторінок описано явними типами.')
    p('Збереження зовнішньої поведінки програми підтверджено незмінним набором із 38 тестів, які '
      'виконувалися після кожного перетворення, а також протоколом роботи програми. Отже, систематичне '
      'автоматизоване тестування є необхідною умовою безпечного рефакторингу, а сам рефакторинг доцільно '
      'виконувати там, де недолік коду ускладнює його розуміння, тестування чи модифікацію.')

    doc.save(OUT)
    print(f'Saved {OUT}')


if __name__ == '__main__':
    build()
